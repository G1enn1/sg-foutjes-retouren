"""Voys Freedom — call-list export importer (gesprekkenlijst, CSV/Excel).

Voys Freedom does not expose an old-style VoIPGRID CDR pull API. The two clean
data routes are (1) the call-list export and (2) the "Gespreksnotificaties"
webhook (see voys_webhook.py). This module handles route (1).

The export has NO employee-name column. Attribution to a person is done via the
internal extension / VoIP account number:
    * outgoing call  -> the employee is the SOURCE   ("Bron")
    * incoming call  -> the employee is the DESTINATION that answered ("Bestemming")
So we match Bron/Bestemming against the set of configured internal extensions
(``voipgrid_user`` in the employee mapping).

Documented export columns (help.voys.nl/export-gesprekken):
    Relatiecode, Klant, Account/Telefoonnummer, Datum, Inkomend/Uitgaand,
    Aantal, Duur, Bron, Bestemming, Bestemmingscode, Wachttijd

Column names/locale can differ per export; adjust ``COLUMNS`` if needed. Verify
against a real export with:  python -m team_monitor.sources.voys_export <file.csv>
"""

from __future__ import annotations

import csv
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from ..config import Settings
from ..models import Channel, Direction, Interaction

# Map our logical field -> possible column headers in the export (case-insensitive).
# Voys Freedom's actual export uses: Datum, Inkomend / Uitgaand, Tijdsduur (in
# SECONDS), Beller (source), Bestemming (destination, "2xx/Naam" for internal).
COLUMNS: dict[str, tuple[str, ...]] = {
    "datum": ("Datum", "Date", "Starttijd", "Start"),
    "direction": ("Inkomend / Uitgaand", "Inkomend/Uitgaand", "Richting", "Direction", "In/Uit"),
    "duur": ("Tijdsduur", "Duur", "Gespreksduur", "Duration", "Talk time"),
    "wachttijd": ("Wachtlijst", "Wachttijd", "Wait time"),
    "bron": ("Beller", "Bron", "Source", "Van", "Caller"),
    "bestemming": ("Bestemming", "Destination", "Naar", "Callee"),
}


def parse_export(path: str | Path, settings: Settings) -> list[Interaction]:
    """Read a Voys Freedom call-list export and return call Interactions."""
    path = Path(path)
    rows = _read_rows(path)
    if not rows:
        return []

    header = {k.strip().lower(): k for k in rows[0].keys()}
    col = {logical: _pick(header, names) for logical, names in COLUMNS.items()}

    ext_index = settings.voipgrid_index()  # internal extension -> Employee
    interactions: list[Interaction] = []

    for raw in rows:
        direction = _direction(_get(raw, col["direction"]))
        bron = _digits(_get(raw, col["bron"]))
        bestemming = _digits(_get(raw, col["bestemming"]))

        # pick the leg that is one of our internal extensions
        emp = None
        if direction is Direction.OUTBOUND and bron in ext_index:
            emp = ext_index[bron]
        elif direction is Direction.INBOUND and bestemming in ext_index:
            emp = ext_index[bestemming]
        else:  # fall back: whichever leg matches a known extension
            emp = ext_index.get(bron) or ext_index.get(bestemming)
        if emp is None:
            continue  # external-to-external or unmapped extension

        interactions.append(
            Interaction(
                employee_id=emp.id,
                employee_name=emp.name,
                channel=Channel.CALL,
                direction=direction,
                timestamp=_parse_dt(_get(raw, col["datum"])),
                duration_seconds=_parse_duration(_get(raw, col["duur"])),
            )
        )
    return interactions


def _read_rows(path: Path) -> list[dict]:
    if path.suffix.lower() in (".xlsx", ".xls"):
        return _read_xlsx(path)
    # CSV — Voys exports are usually ';'-separated (NL locale); sniff to be safe.
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    sample = text[:2048]
    delimiter = ";" if sample.count(";") >= sample.count(",") else ","
    return list(csv.DictReader(text.splitlines(), delimiter=delimiter))


def _read_xlsx(path: Path) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("Excel-export inlezen vereist openpyxl (pip install openpyxl)") from exc
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    headers = [str(h) if h is not None else "" for h in next(rows)]
    return [dict(zip(headers, [("" if v is None else v) for v in r])) for r in rows]


def _pick(header: dict[str, str], candidates: tuple[str, ...]) -> str | None:
    for c in candidates:
        if c.lower() in header:
            return header[c.lower()]
    return None


def _get(raw: dict, key: str | None) -> str:
    if not key:
        return ""
    return str(raw.get(key, "")).strip()


def _digits(value: str) -> str:
    return "".join(ch for ch in value if ch.isdigit())


def _direction(value: str) -> Direction:
    v = value.lower()
    if v.startswith("uit") or "out" in v:
        return Direction.OUTBOUND
    if v.startswith("in"):
        return Direction.INBOUND
    return Direction.INTERNAL


def _parse_duration(value: str) -> int:
    """Duration as HH:MM:SS, MM:SS, or plain seconds -> seconds."""
    value = value.strip()
    if not value:
        return 0
    if ":" in value:
        parts = [int(p) for p in value.split(":") if p.isdigit() or p == "0"]
        seconds = 0
        for p in parts:
            seconds = seconds * 60 + p
        return seconds
    try:
        return int(float(value))
    except ValueError:
        return 0


def _parse_dt(value: str) -> datetime:
    value = value.strip()
    for fmt in ("%d-%m-%Y %H:%M:%S", "%d-%m-%Y %H:%M", "%Y-%m-%d %H:%M:%S",
                "%Y-%m-%dT%H:%M:%S", "%d-%m-%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(value, fmt)
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(value.replace("Z", "+00:00"))
    except ValueError:
        return datetime.now(timezone.utc)


if __name__ == "__main__":  # pragma: no cover - manual check
    import sys

    from ..config import load_settings

    if len(sys.argv) < 2:
        print("Gebruik: python -m team_monitor.sources.voys_export <export.csv> [config.yaml]")
        raise SystemExit(1)
    cfg = load_settings(sys.argv[2] if len(sys.argv) > 2 else None)
    items = parse_export(sys.argv[1], cfg)
    print(f"{len(items)} gesprekken toegewezen aan medewerkers")
    for it in items[:5]:
        print(f"  {it.timestamp:%Y-%m-%d %H:%M}  {it.employee_name:20s}  {it.direction.value:8s}  {it.minutes} min")
