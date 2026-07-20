"""Excel export (optional — requires openpyxl).

Produces one workbook with three sheets: 'Per dag', 'Samenvatting' and
'Moeilijkheid'. If openpyxl is not installed we raise a clear message rather
than failing obscurely.
"""

from __future__ import annotations

from pathlib import Path

from ..aggregate import DailyStat, per_employee_totals
from ..metrics import difficulty_profiles


def write_workbook(stats: list[DailyStat], path: str | Path) -> Path:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError(
            "Excel-export vereist openpyxl. Installeer met: pip install openpyxl"
        ) from exc

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    # --- Sheet 1: per day ---
    ws = wb.active
    ws.title = "Per dag"
    headers = ["Datum", "Medewerker", "E-mails verzonden", "Telefoontjes", "Belminuten",
               "Gem. min/gesprek", "Escalaties", "FCR-ratio", "Gem. CSAT", "Tickets"]
    ws.append(headers)
    for s in stats:
        ws.append([s.day.isoformat(), s.employee_name, s.emails_sent, s.calls_handled,
                   s.call_minutes, s.avg_call_minutes, s.escalations, s.fcr_rate,
                   s.avg_csat, s.ticket_count])

    # --- Sheet 2: per-employee totals ---
    ws2 = wb.create_sheet("Samenvatting")
    ws2.append(["Medewerker", "E-mails", "Telefoontjes", "Belminuten",
                "Gem. min/gesprek", "Escalatie-ratio", "FCR-ratio", "Gem. CSAT"])
    for t in per_employee_totals(stats):
        ws2.append([t.employee_name, t.emails_sent, t.calls_handled, t.call_minutes,
                    t.avg_call_minutes, t.escalation_rate, t.fcr_rate, t.avg_csat])

    # --- Sheet 3: difficulty ---
    ws3 = wb.create_sheet("Moeilijkheid")
    ws3.append(["Medewerker", "Gewogen zwaarte", "Touches/ticket", "Escalatie-ratio",
                "FCR-ratio", "Reopen-ratio", "Gem. CSAT", "Categorie-mix"])
    for p in difficulty_profiles(stats):
        mix = ", ".join(f"{k}:{v}" for k, v in sorted(p.category_mix.items()))
        ws3.append([p.employee_name, p.weighted_difficulty, p.avg_touches_per_ticket,
                    p.escalation_rate, p.fcr_rate, p.reopen_rate, p.avg_csat, mix])

    for sheet in wb.worksheets:
        for col in range(1, sheet.max_column + 1):
            letter = get_column_letter(col)
            sheet[f"{letter}1"].font = Font(bold=True)
            width = max(12, *(len(str(sheet.cell(r, col).value or "")) for r in range(1, min(sheet.max_row, 40) + 1)))
            sheet.column_dimensions[letter].width = min(width + 2, 40)

    wb.save(path)
    return path
